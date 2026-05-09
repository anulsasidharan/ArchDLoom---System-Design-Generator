"""LLD Word document using shared Jinja preamble + structured sections."""

from __future__ import annotations

from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from app.domain.components import ComponentSelectionResult
from app.domain.requirements import ParsedRequirement
from app.services.jinja_env import render_template


def generate_lld_docx(
    requirements: ParsedRequirement,
    components: ComponentSelectionResult,
    *,
    project_title: str | None = None,
    diagram_png: bytes | None = None,
) -> bytes:
    title = project_title or requirements.system_name
    doc = Document()

    h0 = doc.add_heading(f"{title} — Low-Level Design", 0)
    h0.alignment = WD_ALIGN_PARAGRAPH.CENTER

    preamble = render_template(
        "documents/lld_preamble.md.j2",
        project_title=title,
        domain=requirements.domain,
        core_features=requirements.core_features
        or [fr.name for fr in requirements.functional_requirements[:8]],
        nfr=requirements.non_functional_requirements,
        needs_async=requirements.needs_async_processing,
    )
    for block in preamble.split("\n\n"):
        line = block.strip()
        if line:
            doc.add_paragraph(line)

    if diagram_png:
        doc.add_heading("Representative architecture diagram", 1)
        doc.add_picture(BytesIO(diagram_png), width=Inches(6.2))

    doc.add_heading("1. Data model & persistence", 1)
    doc.add_paragraph(
        "Tables below are illustrative placeholders — refine with concrete schemas during sprint 0."
    )
    tbl = doc.add_table(rows=1, cols=4)
    hdr = tbl.rows[0].cells
    for i, label in enumerate(["Table", "Purpose", "Access pattern", "Notes"]):
        hdr[i].text = label
    for row_name in ("users", "sessions", "domain_entities"):
        cells = tbl.add_row().cells
        cells[0].text = row_name
        cells[1].text = "Core transactional entity"
        cells[2].text = "Read/write via application services"
        cells[3].text = "Add indexes after profiling hot queries"

    doc.add_heading("2. Service components", 1)
    for key, decision in sorted(components.selections.items(), key=lambda x: x[0]):
        doc.add_heading(f"{decision.selected.name} ({key})", 2)
        doc.add_paragraph(decision.rationale or "Selected from component library scoring.")
        if decision.selected.features:
            doc.add_paragraph("Notable capabilities:", style="Heading 3")
            for feat in decision.selected.features[:12]:
                doc.add_paragraph(str(feat), style="List Bullet")

    doc.add_heading("3. Inter-service contracts", 1)
    doc.add_paragraph(
        "Expose versioned HTTP or gRPC APIs with backward-compatible deprecations; "
        "document error taxonomy and retry budgets per dependency."
    )

    doc.add_heading("4. Caching & queues", 1)
    if "cache" in components.selections:
        doc.add_paragraph(
            f"Cache tier: {components.selections['cache'].selected.name}. "
            "Define TTL policies per entity volatility."
        )
    else:
        doc.add_paragraph("No dedicated cache tier selected — evaluate hot paths post-MVP.")
    if requirements.needs_async_processing and "message_queue" in components.selections:
        doc.add_paragraph(
            f"Queue: {components.selections['message_queue'].selected.name}. "
            "Specify poison-queue handling and DLQ retention."
        )
    elif requirements.needs_async_processing:
        doc.add_paragraph(
            "Async processing flagged — add an explicit queue selection in future iterations."
        )

    doc.add_heading("5. Reliability & operations", 1)
    nfr = requirements.non_functional_requirements
    bullets = [
        f"Latency target: {nfr.latency_target or 'TBD'}",
        f"Availability target: {nfr.availability_target or 'TBD'}",
        f"RTO/RPO: {nfr.rto or 'TBD'} / {nfr.rpo or 'TBD'}",
    ]
    for b in bullets:
        doc.add_paragraph(b, style="List Bullet")

    doc.add_heading("6. Security hooks", 1)
    if requirements.compliance_needs:
        for c in requirements.compliance_needs:
            doc.add_paragraph(f"{c.name}: {c.description}", style="List Bullet")
    else:
        doc.add_paragraph(
            "Capture authentication, authorization, and encryption decisions "
            "alongside IAM reviews."
        )

    _style_normal(doc)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _style_normal(doc: Document) -> None:
    style = doc.styles["Normal"]
    style.font.size = Pt(11)
