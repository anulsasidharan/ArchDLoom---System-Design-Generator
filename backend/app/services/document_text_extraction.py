"""Extract plain text from uploaded requirement/reference documents."""

from __future__ import annotations

import logging
import re
from io import BytesIO
from pathlib import Path

logger = logging.getLogger(__name__)

# Hard caps to avoid huge prompts and memory use.
MAX_UPLOAD_BYTES = 15 * 1024 * 1024
MAX_CHARS_AFTER_EXTRACT = 120_000

_ALLOWED_SUFFIXES = frozenset({".txt", ".md", ".markdown", ".docx", ".pdf"})


class DocumentExtractError(Exception):
    """Unsupported format, corrupted file, or empty extraction."""


def _truncate(text: str) -> tuple[str, bool]:
    t = text.strip()
    if len(t) <= MAX_CHARS_AFTER_EXTRACT:
        return t, False
    return t[: MAX_CHARS_AFTER_EXTRACT - 80] + "\n\n[… document truncated for processing …]\n", True


def _text_from_docx(data: bytes) -> str:
    from docx import Document

    doc = Document(BytesIO(data))
    parts: list[str] = []
    for p in doc.paragraphs:
        if p.text and p.text.strip():
            parts.append(p.text.strip())
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n\n".join(parts)


def _text_from_pdf(data: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as e:  # pragma: no cover
        raise DocumentExtractError("PDF support is not installed") from e

    reader = PdfReader(BytesIO(data))
    chunks: list[str] = []
    for page in reader.pages:
        try:
            raw = page.extract_text() or ""
        except Exception:
            logger.warning("PDF page text extraction failed", exc_info=True)
            continue
        if raw.strip():
            chunks.append(raw.strip())
    return "\n\n".join(chunks)


def extract_plain_text(filename: str, data: bytes) -> str:
    """Return UTF-8 plain text from supported document bytes."""
    if len(data) > MAX_UPLOAD_BYTES:
        raise DocumentExtractError(
            f"File too large (max {MAX_UPLOAD_BYTES // (1024 * 1024)} MB)",
        )

    suffix = Path(filename or "").suffix.lower()
    if suffix not in _ALLOWED_SUFFIXES:
        raise DocumentExtractError(
            f"Unsupported type {suffix or '(none)'}. Use: {', '.join(sorted(_ALLOWED_SUFFIXES))}",
        )

    try:
        if suffix in (".txt", ".md", ".markdown"):
            text = data.decode("utf-8", errors="replace")
        elif suffix == ".docx":
            text = _text_from_docx(data)
        elif suffix == ".pdf":
            text = _text_from_pdf(data)
        else:
            raise DocumentExtractError(f"Unsupported type {suffix}")
    except DocumentExtractError:
        raise
    except Exception as e:
        logger.exception("Document extraction failed")
        raise DocumentExtractError(f"Could not read document: {e}") from e

    # Normalize excessive whitespace (keeps paragraphs)
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    text, _ = _truncate(text)
    if not text.strip():
        raise DocumentExtractError("No readable text found in the document")
    return text
