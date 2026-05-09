"""High-Level Design Word document: logical architecture and cross-cutting concerns."""

from __future__ import annotations

from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from app.domain.components import ComponentSelectionResult
from app.domain.enterprise import (
    ComplianceMapping,
    CostEstimate,
    HADRStrategy,
    ObservabilityPlan,
    SecurityPlan,
)
from app.domain.requirements import ParsedRequirement
from app.services.jinja_env import render_template


def generate_hld_docx(
    requirements: ParsedRequirement,
    components: ComponentSelectionResult,
    *,
    project_title: str | None = None,
    diagram_png: bytes | None = None,
    network_png: bytes | None = None,
    cost: CostEstimate | None = None,
    compliance_mapping: ComplianceMapping | None = None,
    security: SecurityPlan | None = None,
    hadr: HADRStrategy | None = None,
    observability: ObservabilityPlan | None = None,
) -> bytes:
    title = project_title or requirements.system_name
    doc = Document()

    h0 = doc.add_heading(f"{title} — High-Level Design", 0)
    h0.alignment = WD_ALIGN_PARAGRAPH.CENTER

    total = components.total_monthly_cost_usd
    preamble = render_template(
        "documents/hld_preamble.md.j2",
        project_title=title,
        domain=requirements.domain,
        pattern=components.architecture_pattern.value,
        total_cost=total,
        has_ai=requirements.has_ai_features,
    )
    for block in preamble.split("\n\n"):
        line = block.strip()
        if line:
            doc.add_paragraph(line)

    doc.add_heading("1. System context", 1)
    doc.add_paragraph(
        requirements.summary
        or f"{requirements.system_name} operates in the {requirements.domain} domain "
        f"with {len(requirements.functional_requirements)} major functional areas."
    )
    sm = requirements.scale_metrics
    scale_bits = []
    if sm.daily_active_users is not None:
        scale_bits.append(f"Daily active users (signal): ~{sm.daily_active_users:,}")
    if sm.peak_requests_per_second is not None:
        scale_bits.append(f"Peak RPS (signal): ~{sm.peak_requests_per_second}")
    if sm.data_volume_gb is not None:
        scale_bits.append(f"Data volume (signal): ~{sm.data_volume_gb} GB")
    if scale_bits:
        for b in scale_bits:
            doc.add_paragraph(b, style="List Bullet")
    else:
        doc.add_paragraph(
            "Quantitative scale signals were not specified — validate DAU, peak traffic, "
            "and data growth before production sizing."
        )

    doc.add_heading("2. Logical architecture", 1)
    if diagram_png:
        doc.add_picture(BytesIO(diagram_png), width=Inches(6.2))
        doc.add_paragraph()
    doc.add_paragraph(
        f"The selected pattern is {components.architecture_pattern.value}. "
        "Components below map to deployable or managed services; boundaries align with "
        "failure isolation and team ownership where applicable."
    )

    doc.add_heading("3. Component decisions", 1)
    for key, decision in sorted(components.selections.items(), key=lambda x: x[0]):
        doc.add_heading(f"{decision.selected.name} ({key})", 2)
        p = doc.add_paragraph()
        p.add_run("Category: ").bold = True
        p.add_run(decision.selected.category)
        if decision.selected.vendor:
            p = doc.add_paragraph()
            p.add_run("Vendor: ").bold = True
            p.add_run(decision.selected.vendor)
        doc.add_paragraph(decision.rationale or "Selected via library scoring and requirement fit.")
        if decision.estimated_monthly_cost_usd is not None:
            doc.add_paragraph(
                f"Indicative monthly cost (baseline): ${decision.estimated_monthly_cost_usd:,.0f}"
            )
        if decision.trade_offs.pros:
            doc.add_paragraph("Advantages:", style="Heading 3")
            for pro in decision.trade_offs.pros:
                doc.add_paragraph(pro, style="List Bullet")
        if decision.trade_offs.cons:
            doc.add_paragraph("Limitations:", style="Heading 3")
            for con in decision.trade_offs.cons:
                doc.add_paragraph(con, style="List Bullet")
        if decision.alternatives:
            doc.add_paragraph("Alternatives considered:", style="Heading 3")
            for alt in decision.alternatives[:5]:
                doc.add_paragraph(
                    f"{alt.name}: {alt.rejection_reason}",
                    style="List Bullet",
                )

    doc.add_heading("4. Request and data flows", 1)
    doc.add_paragraph(
        "Typical synchronous path: client → edge/API tier → application services → "
        "data stores and integrations. Authenticate and authorize at the boundary; "
        "propagate correlation IDs for tracing."
    )
    if requirements.needs_async_processing:
        qname = "message_queue"
        if qname in components.selections:
            doc.add_paragraph(
                f"Asynchronous workloads fan out via {components.selections[qname].selected.name}: "
                "publish after transactional commit, consume with idempotent handlers, "
                "and route failures to a dead-letter strategy."
            )
        else:
            doc.add_paragraph(
                "Async processing is required — introduce an explicit queue/bus selection "
                "and document retry and ordering semantics."
            )

    doc.add_heading("5. Interfaces and APIs", 1)
    doc.add_paragraph(
        "Expose stable, versioned HTTP (or gRPC) contracts. Below are thematic surfaces "
        "derived from functional scope — refine paths and schemas during API design review."
    )
    for fr in requirements.functional_requirements[:12]:
        doc.add_paragraph(
            f"{fr.name}: REST or RPC resources backing «{fr.description[:160]}»",
            style="List Bullet",
        )
    if not requirements.functional_requirements:
        doc.add_paragraph(
            "Define resource models and error envelopes alongside the PRD functional themes."
        )

    doc.add_heading("6. Network architecture", 1)
    if network_png:
        doc.add_picture(BytesIO(network_png), width=Inches(6.2))
        doc.add_paragraph()
    doc.add_paragraph(
        "Logical topology: public ingress (WAF + load balancer + NAT) → private app tier → "
        "data tier with default-deny security groups. Egress to external services is allow-listed."
    )

    doc.add_heading("7. Security architecture", 1)
    if security:
        _render_security_section(doc, security)
    else:
        doc.add_paragraph(
            "Encryption in transit: TLS for all external and east-west traffic where supported. "
            "Encryption at rest: enable native KMS-backed encryption on managed databases "
            "and object stores."
        )

    if compliance_mapping and compliance_mapping.frameworks:
        doc.add_heading("8. Compliance mapping", 1)
        _render_compliance_section(doc, compliance_mapping)
        next_index = 9
    else:
        next_index = 8

    doc.add_heading(f"{next_index}. High availability & disaster recovery", 1)
    if hadr:
        _render_hadr_section(doc, hadr)
    else:
        nfr = requirements.non_functional_requirements
        doc.add_paragraph(
            f"Availability target: {nfr.availability_target or 'Align with stakeholder SLOs'}."
        )
        doc.add_paragraph(
            f"Recovery Time Objective (RTO): {nfr.rto or 'TBD'} · "
            f"Recovery Point Objective (RPO): {nfr.rpo or 'TBD'}."
        )
    next_index += 1

    doc.add_heading(f"{next_index}. Monitoring & observability", 1)
    if observability:
        _render_observability_section(doc, observability)
    else:
        doc.add_paragraph(
            "Golden signals: latency, traffic, errors, saturation. Centralize structured logs "
            "with trace correlation; alert on SLO burn rates and dependency health."
        )
    next_index += 1

    doc.add_heading(f"{next_index}. Cost posture", 1)
    if cost:
        _render_cost_section(doc, cost)
    else:
        if total is not None:
            doc.add_paragraph(
                f"Aggregate baseline cost estimate (library): ${total:,.0f}/month — "
                "revisit after load tests."
            )
        else:
            doc.add_paragraph("Cost baseline not computed — refine after sizing exercises.")
    next_index += 1

    doc.add_heading(f"{next_index}. Capacity & scaling", 1)
    nfr = requirements.non_functional_requirements
    doc.add_paragraph(
        f"Throughput posture: {nfr.throughput_target or 'Define peak and sustained QPS'}. "
        f"Concurrency: {nfr.concurrent_users or 'Model concurrent sessions vs. connection pools'}."
    )

    _style_normal(doc)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _render_security_section(doc: Document, security: SecurityPlan) -> None:
    families: list[tuple[str, list[str]]] = [
        ("Authentication", security.authentication),
        ("Authorization", security.authorization),
        ("IAM principles", security.iam_principles),
        ("Encryption — at rest", security.encryption_at_rest),
        ("Encryption — in transit", security.encryption_in_transit),
        ("Network security", security.network_security),
        ("Secrets management", security.secrets_management),
        ("Audit logging", security.audit_logging),
        ("Application security", security.application_security),
        ("Data protection", security.data_protection),
        ("Threat-model notes", security.threat_model_notes),
    ]
    for label, items in families:
        if not items:
            continue
        doc.add_heading(label, 2)
        for item in items:
            doc.add_paragraph(item, style="List Bullet")


def _render_compliance_section(doc: Document, mapping: ComplianceMapping) -> None:
    if mapping.explicit_codes:
        doc.add_paragraph(
            f"Explicitly required: {', '.join(mapping.explicit_codes)}."
        )
    if mapping.inferred_codes:
        doc.add_paragraph(
            f"Inferred from domain or signals: {', '.join(mapping.inferred_codes)}."
        )
    for fw in mapping.frameworks:
        doc.add_heading(fw.name, 2)
        if fw.drivers:
            doc.add_paragraph("Drivers:", style="Heading 3")
            for d in fw.drivers:
                doc.add_paragraph(d, style="List Bullet")
        if fw.controls:
            doc.add_paragraph("Recommended controls:", style="Heading 3")
            tbl = doc.add_table(rows=1, cols=2)
            hdr = tbl.rows[0].cells
            hdr[0].text = "Control"
            hdr[1].text = "Description"
            for ctrl in fw.controls:
                row = tbl.add_row().cells
                row[0].text = ctrl.name
                row[1].text = ctrl.description
    if mapping.notes:
        doc.add_heading("Posture notes", 2)
        for n in mapping.notes:
            doc.add_paragraph(n, style="List Bullet")


def _render_hadr_section(doc: Document, hadr: HADRStrategy) -> None:
    p = doc.add_paragraph()
    p.add_run("Availability target: ").bold = True
    p.add_run(hadr.availability_target)
    p = doc.add_paragraph()
    p.add_run("RTO: ").bold = True
    p.add_run(hadr.rto)
    p.add_run("    ")
    p.add_run("RPO: ").bold = True
    p.add_run(hadr.rpo)
    p = doc.add_paragraph()
    p.add_run("Recommended posture: ").bold = True
    p.add_run(hadr.posture)
    if hadr.posture_rationale:
        doc.add_paragraph(f"Rationale: {hadr.posture_rationale}")
    sections: list[tuple[str, list[str]]] = [
        ("Recommendations", hadr.recommendations),
        ("Backup strategy", hadr.backup_strategy),
        ("Drill cadence", hadr.drills),
    ]
    for label, items in sections:
        if not items:
            continue
        doc.add_heading(label, 2)
        for item in items:
            doc.add_paragraph(item, style="List Bullet")
    if hadr.failover_runbook:
        doc.add_heading("Failover runbook (outline)", 2)
        for i, step in enumerate(hadr.failover_runbook, 1):
            doc.add_paragraph(f"{i}. {step}")


def _render_observability_section(doc: Document, plan: ObservabilityPlan) -> None:
    if plan.pillars:
        doc.add_heading("Pillars", 2)
        for pillar in plan.pillars:
            doc.add_heading(pillar.name, 3)
            doc.add_paragraph(pillar.purpose)
            if pillar.instruments:
                doc.add_paragraph("Instrumentation:", style="Heading 4")
                for instr in pillar.instruments:
                    doc.add_paragraph(instr, style="List Bullet")
            if pillar.alerts:
                doc.add_paragraph("Alerts:", style="Heading 4")
                for alert in pillar.alerts:
                    doc.add_paragraph(alert, style="List Bullet")
    sections: list[tuple[str, list[str]]] = [
        ("SLO targets", plan.slo_targets),
        ("Paging", plan.paging),
        ("Dashboards", plan.dashboards),
    ]
    for label, items in sections:
        if not items:
            continue
        doc.add_heading(label, 2)
        for item in items:
            doc.add_paragraph(item, style="List Bullet")
    if plan.log_retention or plan.trace_sampling:
        doc.add_paragraph(
            f"Log retention: {plan.log_retention or 'TBD'}. "
            f"Trace sampling: {plan.trace_sampling or 'TBD'}."
        )


def _render_cost_section(doc: Document, cost: CostEstimate) -> None:
    doc.add_paragraph(
        f"Aggregate scale-adjusted estimate: ${cost.aggregate_monthly_usd:,.0f} per month "
        f"(pattern: {cost.pattern})."
    )
    if cost.line_items:
        doc.add_heading("Component cost lines", 2)
        tbl = doc.add_table(rows=1, cols=6)
        hdr = tbl.rows[0].cells
        for i, label in enumerate(["Category", "Component", "Baseline", "Scale", "Estimated", "Notes"]):
            hdr[i].text = label
        for li in cost.line_items:
            row = tbl.add_row().cells
            row[0].text = li.category
            row[1].text = li.component
            row[2].text = f"${li.base_cost_usd:,.0f}"
            row[3].text = f"{li.scale_multiplier:.2f}"
            row[4].text = f"${li.estimated_monthly_cost_usd:,.0f}"
            row[5].text = li.notes
    if cost.phases:
        doc.add_heading("Phase projections", 2)
        ptbl = doc.add_table(rows=1, cols=3)
        hdr = ptbl.rows[0].cells
        hdr[0].text = "Phase"
        hdr[1].text = "Monthly cost"
        hdr[2].text = "Notes"
        for phase in cost.phases:
            row = ptbl.add_row().cells
            row[0].text = phase.name
            row[1].text = f"${phase.monthly_cost_usd:,.0f}"
            row[2].text = phase.notes
    if cost.assumptions:
        doc.add_heading("Assumptions", 2)
        for a in cost.assumptions:
            doc.add_paragraph(a, style="List Bullet")


def _style_normal(doc: Document) -> None:
    style = doc.styles["Normal"]
    style.font.size = Pt(11)
