"""Phase 2 PRD-only Word document generation."""

from __future__ import annotations

from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from app.domain.components import ComponentSelectionResult
from app.domain.requirements import ParsedRequirement, SuccessMetric


def generate_prd_docx(
    requirements: ParsedRequirement,
    components: ComponentSelectionResult,
    *,
    project_title: str | None = None,
) -> bytes:
    """Build a PRD .docx from parsed requirements and component decisions."""
    title = project_title or requirements.system_name
    doc = Document()

    h0 = doc.add_heading(f"{title} — Product Requirements Document", 0)
    h0.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_heading("Executive Summary", 1)
    doc.add_paragraph(requirements.summary or _default_summary(requirements))

    doc.add_heading("1. Business Goals", 1)
    for g in requirements.business_goals or ["Deliver core capabilities described below."]:
        doc.add_paragraph(g, style="List Bullet")

    doc.add_heading("2. Success Metrics", 1)
    table = doc.add_table(rows=1, cols=3)
    hdr = table.rows[0].cells
    hdr[0].text = "Metric"
    hdr[1].text = "Target"
    hdr[2].text = "Measurement"
    for cell in hdr:
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True
    metrics = requirements.success_metrics or [
        SuccessMetric(
            name="Availability",
            target=requirements.non_functional_requirements.availability_target or "TBD",
            measurement="Operational dashboard / SLO reviews",
        ),
    ]
    for m in metrics:
        row = table.add_row().cells
        row[0].text = m.name
        row[1].text = m.target
        row[2].text = m.measurement

    doc.add_heading("3. Functional Requirements", 1)
    for i, fr in enumerate(requirements.functional_requirements, start=1):
        doc.add_heading(f"3.{i} {fr.name}", 2)
        doc.add_paragraph(fr.description)
        if fr.user_stories:
            doc.add_paragraph("User stories:", style="Heading 3")
            for us in fr.user_stories:
                doc.add_paragraph(
                    f"As a {us.role}, I want {us.action} so that {us.benefit}.",
                    style="List Bullet",
                )

    doc.add_heading("4. Non-Functional Requirements", 1)
    nfr = requirements.non_functional_requirements
    p = doc.add_paragraph()
    p.add_run("Latency: ").bold = True
    p.add_run(nfr.latency_target or "Not specified")
    p = doc.add_paragraph()
    p.add_run("Availability: ").bold = True
    p.add_run(nfr.availability_target or "Not specified")
    p = doc.add_paragraph()
    p.add_run("Throughput: ").bold = True
    p.add_run(nfr.throughput_target or "Not specified")

    doc.add_heading("5. Compliance & Security", 1)
    if requirements.compliance_needs:
        for c in requirements.compliance_needs:
            doc.add_paragraph(f"{c.name}: {c.description}", style="List Bullet")
    else:
        doc.add_paragraph("None explicitly stated; validate against domain policy.")

    doc.add_heading("6. Technical Constraints & Preferences", 1)
    if requirements.constraints:
        for c in requirements.constraints:
            doc.add_paragraph(c, style="List Bullet")
    else:
        doc.add_paragraph("See tech preferences below.")
    if requirements.tech_preferences:
        doc.add_paragraph("Technology preferences:", style="Heading 3")
        for k, v in requirements.tech_preferences.items():
            doc.add_paragraph(f"{k}: {v}", style="List Bullet")

    doc.add_heading("7. Proposed Architecture Snapshot (Phase 2)", 1)
    if components.total_monthly_cost_usd is not None:
        snap = (
            f"Pattern: {components.architecture_pattern.value}. "
            f"Estimated monthly component cost (library baseline): "
            f"${components.total_monthly_cost_usd:,.0f}"
        )
    else:
        snap = f"Pattern: {components.architecture_pattern.value}."
    doc.add_paragraph(snap)
    for key, decision in sorted(components.selections.items(), key=lambda x: x[0]):
        doc.add_heading(f"{decision.selected.name} ({key})", 2)
        doc.add_paragraph(decision.rationale)
        if decision.trade_offs.pros:
            doc.add_paragraph("Advantages:", style="Heading 3")
            for pro in decision.trade_offs.pros:
                doc.add_paragraph(pro, style="List Bullet")
        if decision.trade_offs.cons:
            doc.add_paragraph("Limitations:", style="Heading 3")
            for con in decision.trade_offs.cons:
                doc.add_paragraph(con, style="List Bullet")

    doc.add_heading("8. Risks & Open Questions", 1)
    for q in requirements.clarification_questions:
        doc.add_paragraph(q, style="List Bullet")
    if not requirements.clarification_questions:
        doc.add_paragraph("Refine scale, SLOs, and compliance scope with stakeholders.")

    _style_normal(doc)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _default_summary(req: ParsedRequirement) -> str:
    return (
        f"{req.system_name} targets the {req.domain} domain with "
        f"{len(req.functional_requirements)} functional themes."
    )


def _style_normal(doc: Document) -> None:
    style = doc.styles["Normal"]
    style.font.size = Pt(11)
