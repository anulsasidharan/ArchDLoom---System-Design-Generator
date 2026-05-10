"""Tests for uploaded document text extraction."""

from __future__ import annotations

from io import BytesIO

import pytest
from docx import Document

from app.services.document_text_extraction import (
    DocumentExtractError,
    extract_plain_text,
)


def test_extract_markdown() -> None:
    raw = "# Title\n\nHello **world**."
    out = extract_plain_text("notes.md", raw.encode("utf-8"))
    assert "Hello" in out


def test_rejects_oversized() -> None:
    with pytest.raises(DocumentExtractError, match="too large"):
        extract_plain_text("huge.txt", b"x" * (16 * 1024 * 1024))


def test_rejects_bad_extension() -> None:
    with pytest.raises(DocumentExtractError, match="Unsupported"):
        extract_plain_text("x.exe", b"abc")


def test_extract_docx() -> None:
    buf = BytesIO()
    d = Document()
    d.add_paragraph("Payment service integration")
    d.add_paragraph("Second line")
    d.save(buf)
    text = extract_plain_text("spec.docx", buf.getvalue())
    assert "Payment service integration" in text
    assert "Second line" in text
